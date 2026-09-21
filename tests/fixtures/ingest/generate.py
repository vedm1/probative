"""Regenerates the synthetic PB1 fixture corpus. Not run by pytest — a
developer action, per S3 ("Fixtures are committed. Recording is a
developer action, replay is the default.").

Needs the dev dependency group (fpdf2, pypdf) in addition to the runtime
extraction libraries (openpyxl, python-docx).

    uv run python tests/fixtures/ingest/generate.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from docx import Document
from fpdf import FPDF
from PIL import Image
from pypdf import PdfReader, PdfWriter

HERE = Path(__file__).parent


def write_plain_txt() -> None:
    (HERE / "plain.txt").write_text(
        "Probative ingests ordinary documents.\n"
        "Every extractor produces a normalised text projection.\n"
        "Offsets into that projection must survive re-extraction.\n",
        encoding="utf-8",
        newline="\n",
    )


def write_crlf_txt() -> None:
    content = "First line.\r\nSecond line.\r\nThird line.\r\n"
    (HERE / "crlf.txt").write_bytes(content.encode("utf-8"))


def write_empty_txt() -> None:
    (HERE / "empty.txt").write_bytes(b"")


def write_notes_md() -> None:
    (HERE / "notes.md").write_text(
        "# Introduction\n"
        "\n"
        "This document introduces the system.\n"
        "\n"
        "## Scope\n"
        "\n"
        "The scope is deliberately narrow.\n"
        "\n"
        "```python\n"
        "# this is not a heading, it is inside a fence\n"
        "print('hello')\n"
        "```\n"
        "\n"
        "### Out of scope\n"
        "\n"
        "Images are out of scope.\n"
        "\n"
        "## Glossary\n"
        "\n"
        "Terms used throughout.\n",
        encoding="utf-8",
        newline="\n",
    )


def write_data_csv() -> None:
    (HERE / "data.csv").write_text(
        "name,role,team\n"
        "Ada,Engineer,Platform\n"
        "Grace,Engineer,Platform\n"
        "Alan,Designer,Product\n"
        "Margaret,PM,Product\n"
        "Katherine,Engineer,Data\n",
        encoding="utf-8",
        newline="\n",
    )


def write_bad_columns_csv() -> None:
    (HERE / "bad_columns.csv").write_text(
        "name,role,team\nAda,Engineer,Platform\nGrace,Engineer\n",
        encoding="utf-8",
        newline="\n",
    )


def write_sheet_xlsx() -> None:
    workbook = openpyxl.Workbook()
    people = workbook.active
    people.title = "People"
    people["A1"] = "Ada"
    people["C1"] = "Engineer"
    people["B7"] = "Platform"

    empty = workbook.create_sheet("Empty")  # noqa: F841 - deliberately left blank

    notes = workbook.create_sheet("Notes")
    notes["A3"] = "Title row is blank above this — no header assumed."

    workbook.save(HERE / "sheet.xlsx")


def write_report_docx() -> None:
    document = Document()
    document.add_heading("Overview", level=1)
    document.add_paragraph("This report summarises the programme.")
    document.add_heading("Timeline", level=2)
    document.add_paragraph("The timeline spans six months.")
    document.add_heading("Risks", level=1)  # pops the "Timeline" H2 back off the stack
    document.add_paragraph("The main risk is scope creep.")
    document.save(HERE / "report.docx")


def _pdf_with_text() -> FPDF:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    pdf.add_page()
    pdf.cell(text="Page one, line one.", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(text="Page one, line two.", new_x="LMARGIN", new_y="NEXT")
    pdf.add_page()
    pdf.cell(text="Page two, line one.", new_x="LMARGIN", new_y="NEXT")
    return pdf


def write_memo_pdf() -> None:
    _pdf_with_text().output(str(HERE / "memo.pdf"))


def write_scanned_pdf() -> None:
    image_path = HERE / "_scan.png"
    Image.new("RGB", (200, 200), color="white").save(image_path)

    pdf = FPDF()
    pdf.add_page()
    pdf.image(str(image_path), x=10, y=10, w=100)
    pdf.output(str(HERE / "scanned.pdf"))
    image_path.unlink()


def write_encrypted_pdf() -> None:
    reader = PdfReader(str(HERE / "memo.pdf"))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password="secret", owner_password="secret-owner")
    with (HERE / "encrypted.pdf").open("wb") as f:
        writer.write(f)


def write_corrupt_pdf() -> None:
    raw = (HERE / "memo.pdf").read_bytes()
    (HERE / "corrupt.pdf").write_bytes(raw[:200])


def main() -> None:
    write_plain_txt()
    write_crlf_txt()
    write_empty_txt()
    write_notes_md()
    write_data_csv()
    write_bad_columns_csv()
    write_sheet_xlsx()
    write_report_docx()
    write_memo_pdf()
    write_scanned_pdf()
    write_encrypted_pdf()
    write_corrupt_pdf()
    print(f"Fixtures written to {HERE}")


if __name__ == "__main__":
    main()
