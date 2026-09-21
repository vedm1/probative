"""XLSX (.xlsx) extraction.

Cell-level, not row/header-based — table semantics are PB2's job. Loaded in
normal (non-read-only) mode deliberately: read-only mode can silently
truncate iteration when a sheet's `<dimension>` metadata undercounts its
actual populated cells, which is exactly the silent-partial-extraction
failure PB1 exists to prevent.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from probative.core.evidence import EmptySourceError, EncryptedSourceError, Locator, LocatorRegion
from probative.ingest._extracted import ExtractedDocument

NORMALIZATION_VERSION = "xlsx/1"


def extract(path: Path) -> ExtractedDocument:
    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except (zipfile.BadZipFile, InvalidFileException) as exc:
        raise EncryptedSourceError(
            path, f"could not open as XLSX — likely password-protected: {exc}"
        ) from exc

    text_parts: list[str] = []
    regions: list[LocatorRegion] = []
    offset = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                value = str(cell.value)
                line = value + "\n"
                start = offset
                end = offset + len(line)
                offset = end
                text_parts.append(line)
                regions.append(
                    LocatorRegion(
                        start=start,
                        end=end,
                        locator=Locator(sheet=sheet.title, cell=cell.coordinate),
                    )
                )

    if not text_parts:
        raise EmptySourceError(path, "workbook has no populated cells on any sheet")

    return ExtractedDocument(
        text="".join(text_parts),
        locator_regions=regions,
        normalization_version=NORMALIZATION_VERSION,
    )
