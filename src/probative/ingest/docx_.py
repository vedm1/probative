"""DOCX (.docx) extraction.

Paragraphs only — tables are out of scope for PB1 (a documented gap; no
fixture contains one). Heading tracking mirrors markdown.py's stack
algorithm, keyed off the paragraph's built-in "Heading N" style.
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.opc.exceptions import PackageNotFoundError

from probative.core.evidence import CorruptSourceError, EncryptedSourceError, Locator, LocatorRegion
from probative.ingest._extracted import ExtractedDocument

NORMALIZATION_VERSION = "docx/1"

_HEADING_STYLE = re.compile(r"^Heading (\d+)$")


def extract(path: Path) -> ExtractedDocument:
    try:
        document = docx.Document(str(path))
    except PackageNotFoundError as exc:
        raise EncryptedSourceError(
            path, f"could not open as DOCX — likely password-protected: {exc}"
        ) from exc
    except Exception as exc:  # pragma: no cover - defensive: any other parse failure
        raise CorruptSourceError(path, f"could not parse as DOCX: {exc}") from exc

    text_parts: list[str] = []
    regions: list[LocatorRegion] = []
    offset = 0
    stack: list[tuple[int, str]] = []

    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        heading_match = _HEADING_STYLE.match(paragraph.style.name if paragraph.style else "")
        if heading_match:
            level = int(heading_match.group(1))
            title = paragraph.text.strip()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))

        line = paragraph.text + "\n"
        start = offset
        end = offset + len(line)
        offset = end
        text_parts.append(line)
        heading_path = [title for _, title in stack] or None
        regions.append(
            LocatorRegion(
                start=start,
                end=end,
                locator=Locator(line=paragraph_index, heading_path=heading_path),
            )
        )

    return ExtractedDocument(
        text="".join(text_parts),
        locator_regions=regions,
        normalization_version=NORMALIZATION_VERSION,
    )
